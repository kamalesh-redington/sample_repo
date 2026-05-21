from config.logger import setup_logger

logger = setup_logger(__name__)

from pptx import Presentation

from .base import DocumentExtractor


class PPTExtractor(DocumentExtractor):

    def extract_text(self, file_path: str) -> str:

        try:

            logger.info(f"Starting PowerPoint extraction: {file_path}")

            prs = Presentation(file_path)

            logger.debug(f"Presentation slide count: {len(prs.slides)}")

            text_parts = []

            for i, slide in enumerate(prs.slides):

                logger.debug(f"Processing slide: {i + 1}")

                text_parts.append(f"\n--- Slide {i+1} ---")

                for shape in slide.shapes:

                    if hasattr(shape, "text") and shape.text.strip():

                        text_parts.append(shape.text)

            final_text = "\n".join(text_parts).strip()

            logger.info(f"PPT extraction completed successfully: {file_path}")

            logger.debug(f"Extracted text length: {len(final_text)}")

            return final_text

        except Exception as e:

            logger.exception(f"PPT extraction failed: {file_path}")

            raise
